"""The explainer deck as an editable PowerPoint, and the gate it must pass.

`export` turns the slide shapes (render.Renderer.deck_model) into native PPTX objects: slide backgrounds, rounded
rectangles, text boxes whose runs keep their size, weight, colour and letter-spacing, the h1 as the slide's
title placeholder, SVG figures as transparent pictures with their text as text boxes, and the speaker notes.
Nothing is a picture of text. `gate` reads the file back and refuses a deck that is not that.
"""
import io
import re
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

EMU = 6350  # one CSS px of the 1920x1080 slide, in EMU (the slide is 12192000 x 6858000)
FONTS = {"Noto Sans", "Noto Serif", "Noto Sans Mono"}
FIXED = datetime(2000, 1, 1)


class PptxError(Exception):
    pass


def _emu(v: float) -> int:
    return int(round(v * EMU))


def _alpha(color_el, a: float) -> None:
    if a < 0.999:
        from lxml import etree
        el = etree.SubElement(color_el, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
        el.set("val", str(max(0, min(100000, round(a * 100000)))))


def _rgb(c):
    from pptx.dml.color import RGBColor
    return RGBColor(round(c["r"]), round(c["g"]), round(c["b"]))


def _fill(shape_fill, c) -> None:
    shape_fill.solid()
    shape_fill.fore_color.rgb = _rgb(c)
    _alpha(shape_fill._xPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"), c["a"])


def _text(tf, it, bg=None) -> None:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
    from pptx.util import Pt
    tf.word_wrap = bool(it["wrap"])
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[it["align"]]
    p.line_spacing = Pt(it["lh"] / 2)
    p.space_before = p.space_after = Pt(0)
    for r in it["runs"]:
        if r.get("br"):
            p.add_line_break()
            continue
        run = p.add_run()
        run.text = r["text"]
        f = run.font
        f.size = Pt(round(r["size"] / 2 * 4) / 4)
        f.bold, f.italic, f.name = r["bold"], r["italic"], r["font"]
        if r.get("color"):
            c = r["color"]
            if c["a"] < 0.999 and bg:  # translucent text is its colour mixed into the slide (PowerPoint clips alpha text)
                c = {k: c[k] * c["a"] + bg[k] * (1 - c["a"]) for k in "rgb"} | {"a": 1}
            f.color.rgb = _rgb(c)
            _alpha(run._r.rPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"), c["a"])
        if r.get("spc"):
            run._r.get_or_add_rPr().set("spc", str(round(r["spc"] / 2 * 100)))


def _rect(slide, it) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu
    rounded = it["radius"] > 0.5
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                Emu(_emu(it["x"])), Emu(_emu(it["y"])), Emu(_emu(it["w"])), Emu(_emu(it["h"])))
    if rounded:
        sh.adjustments[0] = min(0.5, it["radius"] / max(1, min(it["w"], it["h"])))
    if it.get("fill") and it["fill"]["a"] > 0:
        _fill(sh.fill, it["fill"])
    else:
        sh.fill.background()
    if it.get("line"):
        sh.line.width = Emu(_emu(it["line"]["w"]))
        _fill(sh.line.fill, it["line"]["color"])
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    style = sh._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:  # the theme's effect (shadow) reference: the box has exactly the fill and line set here
        sh._element.remove(style)
    sh.name = "Box"


def export(model: list[dict], notes: list[str], title: str, author: str, dest: Path, base: Path) -> None:
    """Write the PPTX. `model`: Renderer.deck_model; `base`: the source folder (image paths)."""
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for i, sl in enumerate(model):
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        if sl.get("bg"):
            _fill(slide.background.fill, sl["bg"])
        used_title = False
        for it in sl["items"]:
            if it["t"] == "rect":
                _rect(slide, it)
            elif it["t"] == "text":
                if it.get("title") and not used_title:
                    used_title, ph = True, slide.shapes.title
                    ph.left, ph.top, ph.width, ph.height = (Emu(_emu(it[k])) for k in ("x", "y", "w", "h"))
                    ph.text_frame.clear()
                    _text(ph.text_frame, it, sl.get("bg"))
                else:
                    tb = slide.shapes.add_textbox(*(Emu(_emu(it[k])) for k in ("x", "y", "w", "h")))
                    _text(tb.text_frame, it, sl.get("bg"))
                    tb.name = "Text"
            elif it["t"] in ("svg", "image"):
                data = io.BytesIO(it["png"]) if it["t"] == "svg" else _read_image(it["src"], base)
                pic = slide.shapes.add_picture(data, *(Emu(_emu(it[k])) for k in ("x", "y", "w", "h")))
                pic.name = it["label"]
                pic._element.nvPicPr.cNvPr.set("descr", it["label"])
        if not used_title:
            slide.shapes.title._element.getparent().remove(slide.shapes.title._element)
        slide.notes_slide.notes_text_frame.text = notes[i]
    cp = prs.core_properties
    cp.title, cp.author, cp.last_modified_by = title, author, author
    cp.created = cp.modified = FIXED
    cp.revision = 1
    buf = io.BytesIO()
    prs.save(buf)
    _write_fixed(buf.getvalue(), dest)


def _read_image(src: str, base: Path):
    import base64
    from urllib.parse import unquote, urlparse
    if src.startswith("data:"):
        return io.BytesIO(base64.b64decode(src.split(",", 1)[1]))
    path = Path(unquote(urlparse(src).path))
    if not path.is_file():
        raise PptxError(f"image not found: {src}")
    return io.BytesIO(path.read_bytes())


def _write_fixed(data: bytes, dest: Path) -> None:
    """Same zip, every member stamped 2000-01-01: the same slides give the same bytes."""
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(2000, 1, 1, 0, 0, 0))
            zi.compress_type, zi.external_attr = zipfile.ZIP_DEFLATED, 0o644 << 16
            zout.writestr(zi, zin.read(info.filename))


# ------------------------------------------------------------------ gate

def _words(text: str) -> Counter:
    return Counter(re.sub(r"\s+", " ", text).split())


def slide_texts(path: Path) -> tuple[list[str], list[str]]:
    """(each slide's text, each slide's notes), read back from the file."""
    from pptx import Presentation
    texts, notes = [], []
    for sl in Presentation(str(path)).slides:
        parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        texts.append("\n".join(parts))
        notes.append(sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else "")
    return texts, notes


def gate(path: Path, expected: list[str], pages: int) -> list[str]:
    """Problems with a built PPTX; empty is clean. `expected`: each slide's text as the PDF page shows it.
    Editable means: one slide per PDF page, speaker notes on every slide, the slide's words all present as text
    (not in a picture), no picture that is the slide, only the vendored Noto faces."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(path))
    problems = []
    slides = list(prs.slides)
    if len(slides) != pages or len(expected) != pages:
        problems.append(f"pptx: {len(slides)} slides for {pages} PDF pages")
    area = prs.slide_width * prs.slide_height
    for i, sl in enumerate(slides, 1):
        notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
        if not notes.strip():
            problems.append(f"pptx slide {i}: no speaker notes")
        have, fonts = "", set()
        for sh in sl.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width * sh.height > area * 0.8:
                problems.append(f"pptx slide {i}: a picture covers the slide (text must stay text, not a flattened image)")
            if sh.has_text_frame:
                have += sh.text_frame.text + "\n"
                fonts |= {r.font.name for p in sh.text_frame.paragraphs for r in p.runs}
        if fonts - FONTS:
            problems.append(f"pptx slide {i}: font outside the vendored faces: {', '.join(sorted(map(str, fonts - FONTS)))}")
        if i <= len(expected):
            want, got = _words(expected[i - 1]), _words(have)
            miss = sorted((want - got).elements())[:6]
            extra = sorted(t for t in (got - want).elements() if re.search(r"\w", t))[:6]  # a decorative mark is no text
            if miss or extra:
                problems.append(f"pptx slide {i}: text differs from the page (missing {miss}, extra {extra})")
    return problems
