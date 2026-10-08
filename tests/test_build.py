"""End to end through the pinned Chromium (run `publishing setup` once)."""
import shutil

import pytest

from publishing import pdf
from publishing.build import BuildError, build, check, discover, resolve
from publishing.cli import main
from publishing.config import load


def new(repo, name, fmt):
    folder = repo / "docs" / "notes" / name
    assert main(["new", str(folder), "--format", fmt]) == 0
    return folder


@pytest.mark.parametrize("fmt,pages", [("deck", 3), ("explainer", 4), ("memo", 2), ("document", 2)])
def test_scaffold_builds_and_checks_current(repo, renderer, fmt, pages):
    folder = new(repo, f"topic-{fmt}-v1", fmt)
    cfg = load(folder)
    s = resolve(folder, cfg)
    assert build(s, cfg, renderer) == "built"
    assert pdf.pages(s.pdf) == pages
    assert pdf.fonts(s.pdf) and all(f.startswith("Publishing") for f in pdf.fonts(s.pdf))
    assert build(s, cfg, renderer) == "current"  # same words and pages: left untouched
    assert check(s, cfg, renderer) is None


def test_document_contents_carry_page_numbers(repo, renderer):
    folder = new(repo, "spec-v1", "document")
    (folder / "document.md").write_text("# Spec\n\n## One\n\nText.\n\n## Two {.newpage}\n\nMore.\n")
    cfg = load(folder)
    s = resolve(folder, cfg)
    build(s, cfg, renderer)
    words = pdf.text(s.pdf)
    assert "One 2 Two 3" in words
    assert pdf.pages(s.pdf) == 3


def test_document_numbering_and_contents_depth(repo, renderer):
    folder = new(repo, "spec-v1", "document")
    body = "## One\n\nText.\n\n### Sub\n\nx\n\n## Two {.newpage}\n\n### Sub\n\nMore.\n"
    cfg = load(folder)
    s = resolve(folder, cfg)

    def words(front):
        (folder / "document.md").write_text(f"---\ntitle: Spec\n{front}---\n{body}")
        build(s, cfg, renderer)
        return pdf.text(s.pdf)

    w = words("numbered: true\n")
    assert "1 One" in w and "2.1 Sub" in w and "2 Two" in w
    assert "2.1 Sub" in w.split("1 One")[1] and "2.1 Sub" not in w.split("1 One")[0]
    w = words("numbered: true\ntoc: 1\n")
    assert "1 One 2 2 Two 3" in w and w.count("Sub") == 2  # the contents list the sections only
    w = words("toc: false\n")
    assert "Contents" not in w and "1 One" not in w
    (folder / "document.md").write_text(f"---\ntitle: Spec\ntoc: 9\n---\n{body}")
    with pytest.raises(BuildError, match="toc"):
        build(s, cfg, renderer)
    (folder / "document.md").write_text(f"---\ntitle: Spec\nnumbered: maybe\n---\n{body}")
    with pytest.raises(BuildError, match="numbered"):
        build(s, cfg, renderer)


def test_check_finds_stale_missing_and_orphan(repo, renderer):
    folder = new(repo, "memo-v1", "memo")
    cfg = load(folder)
    s = resolve(folder, cfg)
    assert check(s, cfg, renderer).endswith("missing (build it: publishing build memo-v1)")
    build(s, cfg, renderer)
    with open(folder / "memo.md", "a") as f:
        f.write("\nOne more line.\n")
    assert "stale" in check(s, cfg, renderer)
    shutil.copy(s.pdf, repo / "docs" / "notes" / "lost-v1.pdf")
    _, orphans = discover(repo / "docs", cfg)
    assert [o.name for o in orphans] == ["lost-v1.pdf"]


def test_lint_stops_a_deck_that_overflows(repo, renderer):
    folder = new(repo, "big-v1", "deck")
    (folder / "slides.py").write_text(
        "from publishing.page import Deck\nd = Deck(__file__, 'big-v1')\n"
        "d.slide('K', 'H', '', '<div style=\"height:2000px\">tall</div>')\nTITLE, S = d.TITLE, d.S\n")
    cfg = load(folder)
    s = resolve(folder, cfg)
    with pytest.raises(BuildError, match="runs into the footer"):
        build(s, cfg, renderer)
    assert not s.pdf.exists()


def test_scan_stops_a_memo_with_a_secret_or_private_word(repo, renderer):
    (repo / "docs" / "report.toml").write_text('project = "Spire"\n[scan]\nwords = ["agent-binding"]\n')
    folder = new(repo, "leak-v1", "memo")
    (folder / "memo.md").write_text("# Leak\n\nThe agent-binding token ghp_" + "a" * 36 + ".\n")
    cfg = load(folder)
    with pytest.raises(BuildError) as e:
        build(resolve(folder, cfg), cfg, renderer)
    assert "secret" in str(e.value) and "project word" in str(e.value)


def test_wide_code_wraps_but_a_wide_table_is_flagged(repo, renderer):
    folder = new(repo, "wide-v1", "memo")
    (folder / "memo.md").write_text("# Wide\n\n```\n" + "x" * 400 + "\n```\n\n"
                                    '<table style="width:1200px"><tr><td>wide</td></tr></table>\n')
    cfg = load(folder)
    with pytest.raises(BuildError, match="wider than the text column"):
        build(resolve(folder, cfg), cfg, renderer)


def test_fixed_document_path_and_markdown_repos(repo, renderer):
    (repo / "docs" / "spire.md").write_text("# Spire\n\n*status: normative*\n\n## Rules\n\nWe hold these today.\n")
    (repo / "docs" / "report.toml").write_text(
        '[[document]]\nsource = "docs/spire.md"\npdf = "docs/Spire.pdf"\npaper = "a4"\ndays = false\n')
    cfg = load(repo)
    s = resolve(repo / "docs" / "spire.md", cfg)
    assert s.pdf.name == "Spire.pdf" and s.paper == "a4"
    assert build(s, cfg, renderer) == "built"
    assert resolve(repo / "docs" / "Spire.pdf", cfg).src == s.src
    sources, orphans = discover(repo / "docs", cfg)
    assert [x.pdf.name for x in sources] == ["Spire.pdf"] and orphans == []

    (repo / "docs" / "report.toml").write_text('format = "markdown"\n')
    with pytest.raises(BuildError, match="markdown itself"):
        build(resolve(repo / "docs" / "spire.md", load(repo)), load(repo), renderer)


def test_pinned_version_mismatch_is_refused(repo, monkeypatch, capsys):
    (repo / "docs" / "report.toml").write_text('publishing = "9.9.9"\n')
    monkeypatch.setenv("PUBLISHING_PINNED", "1")
    with pytest.raises(SystemExit) as e:
        main(["check", str(repo / "docs")])
    assert e.value.code == 3 and "pins 9.9.9" in capsys.readouterr().err  # the toolchain, not the source


def test_compare_puts_pages_side_by_side(repo, renderer, tmp_path):
    a = new(repo, "a-v1", "memo")
    b = new(repo, "b-v1", "document")
    cfg = load(a)
    for f in (a, b):
        build(resolve(f, cfg), cfg, renderer)
    notes = tmp_path / "notes.md"
    notes.write_text("- First improvement.\n- Second improvement.\n")
    out = tmp_path / "cmp-v1.pdf"
    args = ["compare", str(a) + ".pdf", str(b) + ".pdf", "-o", str(out), "--pair", "1:1", "--pair", "2:2:The second page",
            "--notes", str(notes)]
    assert main(args) == 0
    assert pdf.pages(out) == 3 and "Second improvement." in pdf.text(out) and "The second page" in pdf.text(out)
    with pytest.raises(SystemExit) as e:
        main(args)
    assert e.value.code == 2


@pytest.mark.parametrize("family,ok", [("Noto Sans", True), ("Liberation Serif", True), ("sans-serif", False)])
def test_named_families_use_vendored_faces_and_generic_ones_fail(repo, renderer, family, ok):
    folder = new(repo, "fig-v1", "memo")
    (folder / "fig.svg").write_text(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 40">'
                                    f'<text x="10" y="30" font-family="{family}">Label</text></svg>')
    (folder / "memo.md").write_text("# Fig\n\n![A figure](fig.svg)\n")
    cfg = load(folder)
    s = resolve(folder, cfg)
    if ok:
        assert build(s, cfg, renderer) == "built"
    else:
        with pytest.raises(BuildError, match="host font used"):
            build(s, cfg, renderer)


def test_a_staged_copy_builds_with_the_documents_settings(repo, tmp_path):
    """The pre-commit hook renders the staged spire.md from a temp dir to the repo's docs/Spire.pdf."""
    (repo / "docs" / "report.toml").write_text(
        '[[document]]\nsource = "docs/spire.md"\npdf = "docs/Spire.pdf"\npaper = "a4"\ndays = false\n')
    staged = tmp_path / "staged" / "spire.md"
    staged.parent.mkdir()
    staged.write_text("# Spire\n\n## Rules\n\nWe hold these today.\n")
    assert main(["build", str(staged), "-o", str(repo / "docs" / "Spire.pdf")]) == 0
    assert main(["check", str(repo / "docs" / "Spire.pdf")]) == 1  # docs/spire.md itself is missing
    (repo / "docs" / "spire.md").write_text(staged.read_text())
    assert main(["check", str(repo / "docs" / "Spire.pdf")]) == 0


def test_explainer_needs_speaker_notes(repo, renderer):
    folder = new(repo, "talk-v1", "explainer")
    src = (folder / "explainer.py").read_text()
    (folder / "explainer.py").write_text(src.replace('notes="Say what this deck explains and who it is for, in two sentences."', 'notes=""'))
    cfg = load(folder)
    with pytest.raises(BuildError, match="notes are required"):
        build(resolve(folder, cfg), cfg, renderer)


def test_explainer_lint_refuses_a_title_that_wraps(repo, renderer):
    folder = new(repo, "talk-v1", "explainer")
    src = (folder / "explainer.py").read_text()
    (folder / "explainer.py").write_text(src.replace('"One sentence the audience should keep"', '"' + "word " * 30 + '"'))
    cfg = load(folder)
    with pytest.raises(BuildError, match="title wraps"):
        build(resolve(folder, cfg), cfg, renderer)


def test_explainer_pptx_is_editable_with_notes_and_checks_current(repo, renderer):
    from pptx import Presentation

    from publishing import pptx
    folder = new(repo, "talk-v1", "explainer")
    cfg = load(folder)
    s = resolve(folder, cfg)
    assert build(s, cfg, renderer) == "built"
    assert s.pptx.is_file() and s.pptx.name == "talk-v1.pptx"
    prs = Presentation(str(s.pptx))
    texts, notes = pptx.slide_texts(s.pptx)
    assert len(prs.slides) == pdf.pages(s.pdf) == 4
    assert all(n.strip() for n in notes) and notes[1].endswith("Source: Say where the evidence comes from")
    assert "One sentence the audience should keep" in texts[1] and "Output" in texts[2]  # text, also in the diagram
    assert prs.slides[1].shapes.title.text_frame.text == "One sentence the audience should keep"
    assert all(sh.shape_type != 13 or sh.width * sh.height < prs.slide_width * prs.slide_height * 0.8
               for sl in prs.slides for sh in sl.shapes)
    assert build(s, cfg, renderer) == "current"
    assert check(s, cfg, renderer) is None
    s.pptx.unlink()
    assert "missing" in check(s, cfg, renderer)
    assert build(s, cfg, renderer) == "built" and check(s, cfg, renderer) is None


def test_explainer_check_finds_a_stale_pptx(repo, renderer):
    folder = new(repo, "talk-v1", "explainer")
    cfg = load(folder)
    s = resolve(folder, cfg)
    build(s, cfg, renderer)
    src = (folder / "explainer.py").read_text()
    other = new(repo, "talk-v2", "explainer")
    (other / "explainer.py").write_text(src.replace("Walk the three lines in order", "Walk the lines in the order given"))
    s2 = resolve(other, load(other))
    build(s2, cfg, renderer)
    s.pptx.write_bytes(s2.pptx.read_bytes())  # right PDF, another deck's notes
    assert "stale" in check(s, cfg, renderer)


def test_pptx_gate_refuses_a_flattened_slide_and_missing_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Emu

    from publishing import pptx
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    from PIL import Image
    img = tmp_path / "s.png"
    Image.new("RGB", (64, 36), "white").save(img)
    sl.shapes.add_picture(str(img), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(tmp_path / "flat.pptx")
    problems = pptx.gate(tmp_path / "flat.pptx", ["Some words on the slide"], 1)
    text = "\n".join(problems)
    assert "no speaker notes" in text and "flattened" in text and "text differs" in text
    assert pptx.gate(tmp_path / "flat.pptx", ["x"], 2)[0].startswith("pptx: 1 slides for 2")


def test_publish_copies_the_pptx_with_the_pdf(repo, renderer, tmp_path, monkeypatch):
    from publishing.publish import publish
    monkeypatch.setenv("PUBLISHING_DOCUMENTS", str(tmp_path / "Documents"))
    folder = new(repo, "talk-v1", "explainer")
    cfg = load(folder)
    s = resolve(folder, cfg)
    build(s, cfg, renderer)
    state, target = publish(s, cfg)
    assert state == "published" and target.suffix == ".pdf"
    assert (target.parent / "talk-v1.pptx").read_bytes() == s.pptx.read_bytes()
    assert publish(s, cfg)[0] == "already published"


def test_the_explainer_sample_is_current_and_has_every_slide_kind(renderer):
    from pathlib import Path
    sample = Path(__file__).resolve().parent.parent / "docs" / "samples" / "house-style-explainer-v1"
    cfg = load(sample)
    s = resolve(sample, cfg)
    assert check(s, cfg, renderer) is None
    assert pdf.pages(s.pdf) == 13
    html = (sample / "explainer.py").read_text()
    for kind in ("title", "agenda", "section", "statement", "quote", "diagram", "cards", "story", "map", "heatmap", "contrast", "twocol", "closing"):
        assert f"x.{kind}(" in html

def test_explainer_notes_must_match_the_slides_even_when_set_by_hand(repo, renderer):
    folder = new(repo, "talk-v1", "explainer")
    with open(folder / "explainer.py", "a") as f:
        f.write("\nNOTES = NOTES[:-1]\n")  # a hand-set NOTES, bypassing Explainer's own check
    cfg = load(folder)
    with pytest.raises(BuildError, match="every slide needs speaker notes"):
        build(resolve(folder, cfg), cfg, renderer)
